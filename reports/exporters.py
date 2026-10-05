import re
from datetime import date
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='0F766E')
HEADER_FONT = Font(bold=True, color='FFFFFF')
BOLD = Font(bold=True)


def _filename(report, ext):
    return f'vaxsafe-{slugify(report.title)}-{timezone.localdate():%Y%m%d}.{ext}'


def _sheet_name(title, used):
    name = re.sub(r'[\[\]:*?/\\]', '', title)[:31] or 'Sheet'
    base, n = name, 2
    while name in used:
        suffix = f' ({n})'
        name = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def _autosize(ws):
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 2, 10), 60)


def excel_response(report, filter_text, center, user):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws.append([center.name])
    ws['A1'].font = Font(bold=True, size=14)
    if center.address:
        ws.append([center.address])
    ws.append([report.title])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    ws.append([report.description])
    ws.append(['Filters', ' · '.join(filter_text)])
    ws.append(['Generated', f'{timezone.localtime():%Y-%m-%d %H:%M} by {user}'])
    ws.append([])
    for label, value in report.summary:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = BOLD
    _autosize(ws)

    used = {'Summary'}
    for table in report.tables:
        sheet = wb.create_sheet(_sheet_name(table.title, used))
        sheet.append(table.columns)
        for cell in sheet[1]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(vertical='center')
        for row in table.rows:
            sheet.append(row)
        if table.footer:
            sheet.append(table.footer)
            for cell in sheet[sheet.max_row]:
                cell.font = BOLD
        if not table.rows:
            sheet.append(['No records for the selected filters.'])
        if table.note:
            sheet.append([])
            sheet.append([table.note])
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, date):
                    cell.number_format = 'yyyy-mm-dd'
        sheet.freeze_panes = 'A2'
        _autosize(sheet)

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{_filename(report, "xlsx")}"'
    return response


def _link_callback(uri, rel):
    """Let xhtml2pdf load the center logo from disk."""
    if uri.startswith(settings.MEDIA_URL) or uri.startswith('/' + settings.MEDIA_URL):
        path = Path(settings.MEDIA_ROOT) / uri.split(settings.MEDIA_URL, 1)[1]
        return str(path)
    return uri


def pdf_response(report, filter_text, center, user):
    from xhtml2pdf import pisa

    html = render_to_string('reports/pdf.html', {
        'report': report, 'filter_text': filter_text, 'center': center,
        'user': user, 'generated': timezone.localtime(),
        'logo_path': str(Path(center.logo.path)) if center.logo else '',
    })
    buffer = BytesIO()
    result = pisa.CreatePDF(html, dest=buffer, link_callback=_link_callback, encoding='utf-8')
    if result.err:
        return HttpResponse('Could not create the PDF. Try the Excel export instead.', status=500)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{_filename(report, "pdf")}"'
    return response
